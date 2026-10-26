import io
import re

import pdfplumber
from docx import Document


class CVError(Exception):
    """No se pudo leer el CV."""


# Límite para no saturar el LLM con CVs gigantes
MAX_CV_CHARS = 15000


def _leer_txt(contenido):
    """Lee archivos de texto plano intentando varias codificaciones."""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return contenido.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CVError("No pude decodificar el archivo .txt. ¿Está en un formato raro?")


def _leer_docx(contenido):
    doc = Document(io.BytesIO(contenido))
    partes = [p.text for p in doc.paragraphs]
    # Muchos CVs usan tablas para el diseño: también hay que leerlas
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda.text.strip():
                    partes.append(celda.text)
    return "\n".join(partes)


def _leer_pdf(contenido):
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def _normalizar_texto(texto: str) -> str:
    """Limpia espacios múltiples y saltos de línea excesivos."""
    # Reemplaza múltiples espacios por uno solo
    texto = re.sub(r"[ \t]+", " ", texto)
    # Reemplaza 3+ saltos de línea por solo 2
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def leer_cv(contenido, nombre):
    """Recibe los bytes del archivo y su nombre. Devuelve el texto del CV."""
    if not contenido:
        raise CVError("El archivo está vacío.")

    nombre_lower = nombre.lower()

    try:
        if nombre_lower.endswith(".pdf"):
            texto = _leer_pdf(contenido)
        elif nombre_lower.endswith(".docx") or nombre_lower.endswith(".doc"):
            texto = _leer_docx(contenido)
        elif nombre_lower.endswith(".txt"):
            texto = _leer_txt(contenido)
        else:
            raise CVError(
                f"Formato no soportado: '{nombre}'. Usa PDF, DOCX o TXT."
            )
    except CVError:
        raise
    except Exception as e:
        raise CVError(
            "No pude leer el archivo. ¿Está dañado o protegido con contraseña?"
        ) from e

    texto = _normalizar_texto(texto)

    if not texto:
        raise CVError(
            "No encontré texto en el archivo. Si es un PDF escaneado "
            "(una imagen), necesitaría OCR."
        )

    # Recortar si es demasiado largo para el LLM
    if len(texto) > MAX_CV_CHARS:
        texto = texto[:MAX_CV_CHARS] + "\n\n[...CV truncado por longitud...]"

    return texto